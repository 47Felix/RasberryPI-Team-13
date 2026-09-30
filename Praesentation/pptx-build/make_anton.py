#!/usr/bin/env python3
"""Antons Folien (Block 1 Hardware & Sensorik, Block 5 Node-RED & Fernsteuerung),
ausfuehrlicher als im Hauptdeck (~4-5 Min pro Block), mit Sprechtext in den Notizen.
Nutzt dieselben Helfer/Farben wie make_deck.py (build_pptx.py)."""
import importlib.util
spec = importlib.util.spec_from_file_location("bp", "build_pptx.py")
bp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bp)

from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

P = "anton"
COL = bp.PEOPLE[P]
bp.TOTAL = 14

N = 0
def nxt():
    global N
    N += 1
    return N

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def blank_content(chip_text, eyebrow, title):
    s = bp.add_slide()
    bp.set_bg(s, bp.WHITE)
    y = bp.content_header(s, P, nxt(), chip_text, eyebrow, title)
    return s, y + 0.12

def card(s, x, y, w, h, title, body, size=13, filled=False):
    fill = COL["p"] if filled else COL["light"]
    bp.add_rect(s, x, y, w, h, fill=fill, line=None if filled else COL["p"], line_w=1.0, radius=0.08)
    tc = bp.WHITE if filled else COL["dark"]
    bc = bp.WHITE if filled else bp.FG
    bp.add_text(s, x + 0.25, y + 0.2, w - 0.5, 0.4, title, size=15, color=tc, bold=True)
    bp.add_text(s, x + 0.25, y + 0.68, w - 0.5, h - 0.8, body, size=size, color=bc, line_spacing=1.15)

def step_circle(s, x, y, n):
    sp = bp.add_oval(s, x, y, 0.46, 0.46, fill=COL["p"])
    bp.add_text(s, x, y + 0.08, 0.46, 0.3, str(n), size=14, color=bp.WHITE, bold=True, align=PP_ALIGN.CENTER)

# =====================================================================
# BLOCK 1 – Hardware & Sensorik
# =====================================================================
s = bp.lead_slide(nxt(), P, "Block 1: Hardware & Sensorik", "Anton", eyebrow="Challenge I")
notes(s, "Danke Erik. Ich fange ganz vorne in der Kette an: bei der Hardware und den Sensoren.")

# ---------- Systemarchitektur ----------
s, y = blank_content("Challenge I", "Überblick", "Systemarchitektur")
bp.diagram_box(s, 0.75, y + 0.1, 5.2, 1.9, "Sensor-Board · I2C 0x08",
               ["KY-028 Temperatursensor → A0", "LED (PWM) → D9"], P)
bp.diagram_box(s, 7.38, y + 0.1, 5.2, 1.9, "Aktor-Board · I2C 0x09",
               ["KY-028 Temperatursensor → A0", "LED (PWM) → D5",
                "Lüfter (über Transistor) → D4", "Servo / Ventil → D6"], P)
bus_y = y + 2.55
bp.add_line(s, 3.35, y + 2.0, 3.35, bus_y, COL["p"], width_pt=2)
bp.add_line(s, 9.98, y + 2.0, 9.98, bus_y, COL["p"], width_pt=2)
bp.add_line(s, 3.35, bus_y, 9.98, bus_y, COL["p"], width_pt=2)
bp.add_text(s, 3.35, bus_y - 0.36, 6.63, 0.3, "I2C-Bus · SDA + SCL · Pull-ups auf 3,3 V",
            size=12, color=COL["dark"], bold=True, align=PP_ALIGN.CENTER)
bp.diagram_arrow_v(s, 6.667, bus_y, bus_y + 0.55, COL["p"])
bp.diagram_box(s, 3.9, bus_y + 0.6, 5.53, 1.15, "Raspberry Pi = Master („Gehirn“)",
               ["liest beide Boards · rechnet & regelt · speichert in SQLite"], P, filled=True)
notes(s, "Unser Aufbau besteht aus zwei Arduino Unos und einem Raspberry Pi. Die Arduinos sitzen direkt an "
         "der Hardware, sie messen und steuern. Der Pi ist das Gehirn: Er holt die Werte ab, entscheidet, was "
         "passiert, und speichert alles. Alle drei hängen an einem gemeinsamen I2C-Bus, das sind nur zwei "
         "Leitungen, Daten und Takt. Damit der Pi weiß, mit wem er redet, hat jedes Board eine eigene Adresse: "
         "0x08 ist unser Sensor-Board, 0x09 das Aktor-Board, weil da zusätzlich Lüfter und Ventil dranhängen. "
         "Der Pi ist der Master, er fragt an, die Arduinos antworten.")

# ---------- Verkabelung ----------
s, y = blank_content("Challenge I", "Hardware", "Verkabelung im Überblick")
rows = [
    ["Board", "Bauteil", "Pin"],
    ["Sensor-Board (0x08)", "KY-028 Temperatursensor", "A0 (analog)"],
    ["", "LED", "D9 (PWM)"],
    ["Aktor-Board (0x09)", "KY-028 Temperatursensor", "A0 (analog)"],
    ["", "LED", "D5 (PWM)"],
    ["", "Lüfter über Transistor", "D4"],
    ["", "Servo / Ventil", "D6"],
    ["Beide", "I2C zum Pi", "A4 = SDA, A5 = SCL → Pi GPIO2 / GPIO3"],
]
bp.simple_table(s, 0.75, y, 7.6, rows, [2.2, 2.5, 2.9], P, row_h=0.5, font_size=13)
card(s, 8.75, y, 3.83, 4.0, "⚠ Warum 3,3 V?",
     "Arduino arbeitet mit **5 V**, der Pi nur mit **3,3 V**.\n\n"
     "5 V auf dem Bus könnten die Pi-GPIOs zerstören.\n\n"
     "→ Pull-up-Widerstände auf **3,3 V**\n→ interne 5-V-Pull-ups der Arduinos **abgeschaltet**\n→ gemeinsame Masse (GND)",
     size=13)
notes(s, "Hier seht ihr, was wo angeschlossen ist. Beide Boards haben den gleichen Temperatursensor am "
         "Analogeingang A0 und eine LED an einem PWM-Pin. Am Aktor-Board hängen zusätzlich der Lüfter und der "
         "Servo, das erklärt Dogan gleich genauer. Wichtig bei I2C: Der Arduino arbeitet mit 5 Volt, der Pi nur "
         "mit 3,3 Volt. Kämen 5 Volt auf den Bus, könnten wir den Pi beschädigen. Deshalb ziehen wir die "
         "Leitungen mit Widerständen auf 3,3 Volt hoch und haben die internen 5-Volt-Pull-ups der Arduinos "
         "abgeschaltet. Außerdem haben alle drei Geräte eine gemeinsame Masse.")

# ---------- KY-028 ----------
s, y = blank_content("Challenge I", "Bewertung: Sensor funktionstüchtig (5 Pkt)", "Der Sensor: KY-028")
bp.add_bullets(s, 0.75, y, 6.6, 4.6, [
    "Modul mit **NTC-Thermistor** – ein Widerstand, der sich mit der Temperatur ändert",
    "Wir nutzen den **analogen Ausgang** (Pin A0)",
    "Der ADC des Arduino macht daraus eine Zahl **0–1023** (10 Bit) = **Rohwert**",
    "NTC = „Negative Temperature Coefficient“: **wärmer → Rohwert kleiner**",
    "Messung alle **200 ms**",
    "An den Pi: **2 Bytes** (High- + Low-Byte), weil 1023 nicht in ein Byte (max. 255) passt",
], size=15, bullet_color=COL["p"])
# stat callout
bp.add_rect(s, 7.75, y, 4.83, 4.4, fill=COL["light"], line=None, radius=0.06)
bp.add_text(s, 7.75, y + 0.25, 4.83, 0.35, "Gemessen am Sensor-Board", size=13, color=COL["dark"], bold=True, align=PP_ALIGN.CENTER)
bp.add_text(s, 7.95, y + 0.8, 2.1, 0.5, "23 °C", size=22, color=bp.MUTED, bold=True, align=PP_ALIGN.CENTER, font=bp.HEAD_FONT)
bp.add_text(s, 10.28, y + 0.8, 2.1, 0.5, "30 °C", size=22, color=bp.MUTED, bold=True, align=PP_ALIGN.CENTER, font=bp.HEAD_FONT)
bp.add_text(s, 7.95, y + 1.35, 2.1, 1.0, "212", size=54, color=COL["p"], bold=True, align=PP_ALIGN.CENTER, font=bp.HEAD_FONT)
bp.add_text(s, 10.28, y + 1.35, 2.1, 1.0, "160", size=54, color=COL["p"], bold=True, align=PP_ALIGN.CENTER, font=bp.HEAD_FONT)
bp.add_text(s, 9.7, y + 1.6, 0.9, 0.6, "→", size=32, color=COL["dark"], align=PP_ALIGN.CENTER)
bp.add_text(s, 7.95, y + 2.5, 4.43, 0.35, "Rohwert", size=12, color=bp.MUTED, align=PP_ALIGN.CENTER)
bp.add_text(s, 7.95, y + 3.2, 4.43, 0.9, "wärmer  →  kleinerer Rohwert", size=17, color=COL["dark"],
            bold=True, italic=True, align=PP_ALIGN.CENTER, font=bp.HEAD_FONT)
notes(s, "Als Sensor nutzen wir das KY-028. Darauf sitzt ein NTC-Thermistor, ein Widerstand, dessen Wert sich "
         "mit der Temperatur ändert. Der Arduino misst die Spannung daran und wandelt sie mit seinem "
         "Analog-Digital-Wandler in eine Zahl zwischen 0 und 1023 um, das ist unser Rohwert. NTC heißt: Wird es "
         "wärmer, sinkt der Wert. Bei 23 Grad haben wir zum Beispiel rund 212 gemessen, bei 30 Grad nur noch "
         "160. Der Arduino misst fünfmal pro Sekunde. Weil der Wert bis 1023 gehen kann und ein Byte nur bis "
         "255, schicken wir ihn als zwei Bytes an den Pi.")

# ---------- DHT22 -> KY-028 ----------
s, y = blank_content("Challenge I", "Technische Entscheidung", "Vom DHT22 zum KY-028")
cw, gap = 3.73, 0.3
card(s, 0.75, y + 0.1, cw, 3.2, "1 · Plan",
     "**DHT22** auf dem Aktor-Board\n\nDigitaler Sensor für Temperatur **und** Luftfeuchte")
card(s, 0.75 + cw + gap, y + 0.1, cw, 3.2, "2 · Problem",
     "Nach jedem Reset nur **NaN** – „Not a Number“, also kein gültiger Wert\n\n"
     "Geprüft per seriellem Monitor, Verkabelung mehrfach kontrolliert, I2C lief parallel fehlerfrei")
card(s, 0.75 + 2 * (cw + gap), y + 0.1, cw, 3.2, "3 · Entscheidung",
     "DHT22 raus, **zweites KY-028** rein\n\nSensor war vermutlich defekt – nicht weiter raten, sondern ersetzen",
     filled=True)
bp.add_text(s, 0.75, y + 3.65, 11.83, 0.5,
            "Vorteil: Beide Boards messen jetzt nach **demselben Prinzip** – einheitlicher Code, vergleichbare Werte.",
            size=16, color=bp.FG)
notes(s, "Ursprünglich war das nicht so geplant. Auf dem Aktor-Board sollte ein DHT22 sitzen, der Temperatur und "
         "Luftfeuchte misst. Der hat aber nie einen gültigen Wert geliefert, immer nur NaN, also ‚keine Zahl‘. "
         "Wir haben das direkt über den seriellen Monitor geprüft, damit wir einen Fehler in unserer Software "
         "ausschließen konnten, und die Verkabelung mehrfach kontrolliert. Der Sensor war vermutlich defekt. "
         "Statt weiter Zeit zu verlieren, haben wir ihn durch ein zweites KY-028 ersetzt. Das hat den Vorteil, "
         "dass jetzt beide Boards gleich funktionieren.")

# ---------- Kalibrierung ----------
s, y = blank_content("Challenge I", "Bewertung: Formate der Sensordaten (5 Pkt)", "Kalibrierung: Rohwert → °C")
bp.add_text(s, 0.75, y, 11.8, 0.4,
            "2-Punkt-Kalibrierung mit einem echten **Referenzthermometer**:", size=16)
bp.simple_table(s, 0.75, y + 0.55, 7.2, [
    ["", "Punkt 1", "Punkt 2"],
    ["Sensor-Board", "23,0 °C → 212", "30,0 °C → 160"],
    ["Aktor-Board", "23,0 °C → 174", "30,5 °C → 126"],
], [2.2, 2.5, 2.5], P, row_h=0.55, font_size=14)
bp.add_bullets(s, 0.75, y + 2.5, 7.2, 2.5, [
    "Dazwischen **lineare Interpolation** – eine Gerade durch beide Punkte",
    "Jedes Board hat **eine eigene Kurve** – Bauteile streuen (212 vs. 174 bei gleicher Temperatur!)",
    "Umrechnung macht der **Pi** (`calibration.py`), die Arduinos schicken nur den Rohwert",
], size=15, bullet_color=COL["p"])
# datenfluss-format
bp.add_rect(s, 8.35, y + 0.55, 4.23, 4.3, fill=COL["light"], line=None, radius=0.06)
bp.add_text(s, 8.6, y + 0.72, 3.8, 0.35, "Datenformat entlang der Kette", size=13, color=COL["dark"], bold=True)
steps = [("Sensor", "Spannung (analog)"), ("Arduino", "Rohwert 0–1023"), ("I2C", "2 Bytes"),
         ("Pi", "°C (Kommazahl)"), ("SQLite", "Zeile mit Zeitstempel")]
for i, (a, b) in enumerate(steps):
    yy = y + 1.2 + i * 0.72
    bp.add_text(s, 8.6, yy, 1.2, 0.3, a, size=13, color=COL["p"], bold=True)
    bp.add_text(s, 9.8, yy, 2.7, 0.3, b, size=13, color=bp.FG)
    if i < len(steps) - 1:
        bp.add_text(s, 8.6, yy + 0.33, 1.0, 0.3, "↓", size=13, color=COL["p"])
notes(s, "Ein Rohwert von 212 sagt noch nichts über die Temperatur. Deshalb haben wir kalibriert: Wir haben ein "
         "echtes Thermometer neben den Sensor gelegt und bei zwei Temperaturen den Rohwert notiert. Durch diese "
         "zwei Punkte legen wir eine Gerade und können so jeden Rohwert in Grad Celsius umrechnen. Wie ihr seht, "
         "liefern die beiden Sensoren bei gleicher Temperatur ganz unterschiedliche Werte, 212 gegenüber 174. "
         "Das ist normale Bauteilstreuung, deshalb bekommt jedes Board seine eigene Kurve. Die Umrechnung "
         "passiert auf dem Pi, die Arduinos schicken nur den Rohwert. Rechts seht ihr, wie sich das Datenformat "
         "auf dem Weg ändert: von der Spannung über den Rohwert und zwei Bytes bis zur Temperatur in der Datenbank.")

# ---------- LED / PWM ----------
s, y = blank_content("Challenge I", "Bewertung: LED + Helligkeit folgt Sensor (4 + 6 Pkt)", "LED zeigt den Messwert")
bp.demo_badge(s, 0.75, y)
bp.add_bullets(s, 0.75, y + 0.5, 6.4, 2.6, [
    "**PWM** (Pulsweitenmodulation): LED wird sehr schnell an/aus geschaltet",
    "Anteil „an“ pro Takt = Helligkeit → Wert **0 (aus) bis 255 (voll)**",
    "Rechnung **direkt auf dem Arduino** – kein Umweg über den Pi",
    "→ reagiert sofort und funktioniert auch ohne Pi",
], size=15, bullet_color=COL["p"])
bp.code_block(s, 0.75, y + 3.3, 6.4, 1.3,
              "hell = map(rohwert, ..., 255, 0);\nhell = constrain(hell, 0, 255);\nanalogWrite(PIN_LED, hell);", P)
# PWM duty-cycle illustration
bx, bw = 7.75, 4.83
bp.add_rect(s, bx, y, bw, 4.6, fill=COL["light"], line=None, radius=0.05)
bp.add_text(s, bx + 0.3, y + 0.2, bw - 0.6, 0.35, "PWM: gleicher Takt, andere Einschaltdauer", size=13, color=COL["dark"], bold=True)
for i, (duty, label) in enumerate([(0.15, "kalt → dunkel"), (0.5, "mittel"), (0.9, "warm → hell")]):
    ly = y + 0.95 + i * 1.2
    bp.add_text(s, bx + 0.3, ly, 2.0, 0.3, label, size=12, color=bp.FG, bold=True)
    period = 1.3
    for k in range(3):
        px = bx + 0.3 + k * period
        on_w = period * duty
        bp.add_rect(s, px, ly + 0.35, on_w, 0.45, fill=COL["p"], line=None, shape=MSO_SHAPE.RECTANGLE)
        bp.add_line(s, px + on_w, ly + 0.8, px + period, ly + 0.8, COL["p"], width_pt=2)
notes(s, "Jedes Board hat eine LED, die die Temperatur anzeigt: Je wärmer, desto heller. Dafür nutzen wir PWM. Die "
         "LED wird so schnell an- und ausgeschaltet, dass man das Flackern nicht sieht. Je länger sie pro Takt an "
         "ist, desto heller wirkt sie, das seht ihr rechts. Der Arduino rechnet den Rohwert mit der Funktion "
         "map() direkt in einen Helligkeitswert zwischen 0 und 255 um, constrain() sorgt dafür, dass der Wert "
         "nicht aus diesem Bereich rausläuft. Das passiert komplett auf dem Arduino, ohne den Pi. So reagiert die "
         "LED sofort und funktioniert auch, wenn der Pi mal nicht antwortet. Ich zeige das kurz.\n\n"
         "[DEMO: Finger auf den Sensor am SENSOR-BOARD legen – dort reagiert die LED zwischen ca. 29 und 33 °C, "
         "also gut sichtbar. Warten, bis sie heller wird.]\n\n"
         "Damit gebe ich an Felix ab, der zeigt, wie die Daten zum Pi kommen.")

# =====================================================================
# BLOCK 5 – Node-RED & Fernsteuerung
# =====================================================================
s = bp.lead_slide(nxt(), P, "Block 5: Node-RED & Fernsteuerung", "Anton", eyebrow="Challenge II")
notes(s, "Danke Erik. Ich zeige jetzt, wie Node-RED die Brücke zwischen unserem Pi und MQTT schlägt.")

# ---------- Was ist Node-RED ----------
s, y = blank_content("Challenge II", "Entwicklungswerkzeug", "Was ist Node-RED?")
bp.add_bullets(s, 0.75, y, 6.6, 4.6, [
    "Werkzeug zur **visuellen Programmierung**, läuft auf dem Pi (Port 1880)",
    "Programme = **Flows**, gebaut aus **Nodes** (Bausteinen) und Verbindungen",
    "Nachrichten (`msg`) wandern von Node zu Node",
    "Fertige Nodes für **MQTT, SQLite, Dateien, Skriptaufrufe**",
    "Eigene Logik in **Function-Nodes** (JavaScript)",
    "Bei uns: **die Brücke zwischen Pi-Backend und MQTT – in beide Richtungen**",
], size=15, bullet_color=COL["p"])
nx = 8.1
for i, (t, sub) in enumerate([("Inject", "Auslöser"), ("Function", "eigene Logik"), ("MQTT out", "senden")]):
    ny = y + 0.2 + i * 1.45
    bp.add_rect(s, nx, ny, 3.9, 0.85, fill=COL["light"], line=COL["p"], line_w=1.25, radius=0.3)
    bp.add_text(s, nx + 0.3, ny + 0.14, 3.3, 0.3, t, size=15, color=COL["dark"], bold=True)
    bp.add_text(s, nx + 0.3, ny + 0.46, 3.3, 0.3, sub, size=12, color=bp.MUTED)
    if i < 2:
        bp.diagram_arrow_v(s, nx + 1.95, ny + 0.85, ny + 1.45, COL["p"], label="msg", label_side="right")
notes(s, "Node-RED ist ein Werkzeug, mit dem man Abläufe grafisch zusammenbaut, statt alles von Hand zu "
         "programmieren. Man zieht Bausteine, sogenannte Nodes, auf eine Fläche und verbindet sie. Durch diese "
         "Verbindungen wandern Nachrichten von Baustein zu Baustein, wie rechts gezeigt. Für MQTT und "
         "Datenbanken gibt es fertige Nodes, und wo wir eigene Logik brauchen, schreiben wir sie in einem "
         "Function-Node mit JavaScript. Bei uns ist Node-RED die Brücke zwischen unserem Python-Backend aus "
         "Challenge I und dem MQTT-Broker, und zwar in beide Richtungen.")

# ---------- Flow Überblick ----------
s, y = blank_content("Challenge II", "Bewertung: Steuerung über Node-RED realisiert (9 Pkt)", "Unser Flow: zwei Richtungen")
bp.add_text(s, 0.75, y, 11.8, 0.35, "① Pi → Handy: Messwerte veröffentlichen", size=15, color=COL["dark"], bold=True)
lane1 = [("Inject", "alle 5 s"), ("SQLite", "letzte Messung"), ("Function", "Zeile → 8 Topics"), ("MQTT out", "retained")]
bw_, gap_ = 2.55, 0.52
for i, (t, sub) in enumerate(lane1):
    x = 0.75 + i * (bw_ + gap_)
    bp.diagram_box(s, x, y + 0.5, bw_, 1.0, t, [sub], P)
    if i < 3:
        bp.diagram_arrow_h(s, x + bw_ + 0.04, x + bw_ + gap_ - 0.04, y + 1.0, COL["p"])
bp.add_text(s, 0.75, y + 1.95, 11.8, 0.35, "② Handy → Pi: Befehle empfangen", size=15, color=COL["dark"], bold=True)
bp.diagram_box(s, 0.75, y + 2.6, bw_, 1.0, "MQTT in", ["control/#"], P)
bp.diagram_arrow_h(s, 0.75 + bw_ + 0.04, 0.75 + bw_ + gap_ - 0.04, y + 3.1, COL["p"])
bp.diagram_box(s, 0.75 + bw_ + gap_, y + 2.3, bw_, 1.6, "Function", ["validieren", "(prüfen + begrenzen)"], P, filled=True)
outs = [("Debug", "Anzeige im Editor"), ("Datei", "control_log.ndjson"), ("Exec", "set_control.py")]
ox = 0.75 + 2 * (bw_ + gap_)
for i, (t, sub) in enumerate(outs):
    oy = y + 2.2 + i * 0.62
    bp.add_rect(s, ox, oy, 5.62, 0.52, fill=COL["light"], line=COL["p"], line_w=1.0, radius=0.2)
    bp.add_text(s, ox + 0.2, oy + 0.12, 1.2, 0.3, t, size=13, color=COL["dark"], bold=True)
    bp.add_text(s, ox + 1.4, oy + 0.12, 4.0, 0.3, sub, size=13, color=bp.FG)
    bp.add_line(s, ox - gap_ + 0.04, y + 3.1, ox - 0.02, oy + 0.26, COL["p"], width_pt=1.5)
bp.add_text(s, 0.75, y + 4.3, 11.8, 0.4, "Tipp: Hier ggf. einen Screenshot des echten Flows aus dem Node-RED-Editor einfügen.",
            size=12, color=bp.MUTED, italic=True)
notes(s, "Hier seht ihr unseren Flow. Er hat zwei Stränge. Oben geht es vom Pi zum Handy: Messwerte werden aus "
         "der Datenbank abgeholt und per MQTT veröffentlicht. Unten geht es vom Handy zum Pi: Befehle kommen "
         "rein, werden geprüft und dann an drei Stellen weitergegeben. Beide Richtungen schauen wir uns jetzt "
         "einzeln an.\n\n[Den Hinweis unten auf der Folie vor dem Vortrag löschen oder durch einen Screenshot ersetzen.]")

# ---------- Richtung 1 ----------
s, y = blank_content("Challenge II", "Richtung ①  Pi → Handy", "Messwerte veröffentlichen")
steps1 = [
    ("Inject-Node", "löst **alle 5 Sekunden** aus"),
    ("SQLite-Node", "holt die **neueste Zeile** aus `challenge_i.db`"),
    ("Function-Node", "verteilt sie auf **8 Topics** und berechnet die **Ø-Temperatur**"),
    ("MQTT out", "veröffentlicht alles **retained** am Mosquitto-Broker"),
]
for i, (t, d) in enumerate(steps1):
    yy = y + i * 0.95
    step_circle(s, 0.75, yy, i + 1)
    bp.add_text(s, 1.4, yy + 0.02, 5.5, 0.3, t, size=15, color=COL["dark"], bold=True)
    bp.add_text(s, 1.4, yy + 0.38, 5.6, 0.5, d, size=13.5, color=bp.FG)
bp.code_block(s, 0.75, y + 3.95, 6.3, 0.6, "SELECT * FROM readings ORDER BY id DESC LIMIT 1;", P)
bp.add_rect(s, 7.45, y, 5.13, 4.55, fill=COL["light"], line=None, radius=0.05)
bp.add_text(s, 7.7, y + 0.18, 4.7, 0.35, "team13-1/icetruck/ …", size=14, color=COL["dark"], bold=True)
topics = ["sensors/sensor_board_raw", "sensors/sensor_board_temp_c", "sensors/actor_board_raw",
          "sensors/actor_board_temp_c", "sensors/avg_temp_c", "actuators/fan_pwm",
          "actuators/valve_angle", "status  (alles als JSON)"]
bp.add_bullets(s, 7.7, y + 0.65, 4.7, 3.8, [f"`{t}`" if "JSON" not in t else "`status` (alles als JSON)" for t in topics],
               size=13, bullet_color=COL["p"], space_after=4)
notes(s, "Alle fünf Sekunden startet ein Inject-Node den Ablauf. Danach holt ein SQLite-Node mit dieser "
         "SQL-Abfrage die neueste Zeile aus unserer Datenbank, also die letzte Messung. Ein Function-Node "
         "verteilt sie auf acht einzelne MQTT-Topics: die Temperaturen beider Sensoren, die Rohwerte, die Werte "
         "für Lüfter und Ventil und zusätzlich die Durchschnittstemperatur, die wir direkt hier im Flow "
         "berechnen. Außerdem gibt es ein Status-Topic mit allen Werten als JSON. Alles wird retained "
         "veröffentlicht. Das heißt, der Broker merkt sich den letzten Wert, und ein Handy, das sich neu "
         "verbindet, sieht ihn sofort, statt auf den nächsten Zyklus zu warten.")

# ---------- Richtung 2 ----------
s, y = blank_content("Challenge II", "Richtung ②  Handy → Pi", "Befehle vom Handy prüfen")
bp.add_text(s, 0.75, y, 11.8, 0.4,
            "MQTT-In abonniert `team13-1/icetruck/control/#`  (# = alle Unter-Topics)", size=15)
bp.simple_table(s, 0.75, y + 0.6, 11.83, [
    ["Topic", "Erlaubt", "Bei ungültigem Wert"],
    ["`control/mode/set`", "auto  oder  manual", "verworfen + Warnung"],
    ["`control/fan_pwm/set`", "0 – 255", "auf Grenze begrenzt (999 → 255)"],
    ["`control/valve_angle/set`", "0 – 180", "auf Grenze begrenzt"],
    ["unbekanntes Topic", "–", "verworfen + Warnung"],
], [4.2, 3.2, 4.43], P, row_h=0.52, font_size=14)
bp.add_text(s, 0.75, y + 3.45, 11.8, 0.35, "Gültiger Befehl geht an drei Ausgänge:", size=15, color=COL["dark"], bold=True)
for i, (t, d) in enumerate([("Debug", "sofort im Editor sichtbar"), ("Log-Datei", "mit Zeitstempel – wer hat wann was geschaltet?"),
                            ("Exec → set_control.py", "gibt den Befehl ans Backend weiter")]):
    x = 0.75 + i * 4.03
    bp.add_rect(s, x, y + 3.9, 3.77, 0.95, fill=COL["light"], line=None, radius=0.1)
    bp.add_text(s, x + 0.2, y + 3.98, 3.4, 0.3, t, size=14, color=COL["dark"], bold=True)
    bp.add_text(s, x + 0.2, y + 4.32, 3.4, 0.5, d, size=12, color=bp.FG)
notes(s, "In die andere Richtung abonniert der Flow alle Topics unter control. Die Raute ist ein Platzhalter und "
         "bedeutet ‚alles darunter‘. Jeder Befehl wird zuerst geprüft, denn wir wollen nicht, dass jemand "
         "Unsinn an die Hardware schickt. Beim Modus sind nur auto und manual erlaubt, alles andere wird "
         "verworfen. Zahlen werden auf den erlaubten Bereich begrenzt: Schickt jemand für den Lüfter 999, wird "
         "daraus 255, weil mehr nicht geht, und beim Ventil ist bei 180 Grad Schluss. Ein gültiger Befehl geht "
         "an drei Stellen: in die Debug-Ansicht, in eine Log-Datei mit Zeitstempel, damit wir nachvollziehen "
         "können, wer wann was geschaltet hat, und an ein Python-Skript, das ihn ans Backend weitergibt.")

# ---------- Auto / Manuell ----------
s, y = blank_content("Challenge II", "Bewertung: Zusätzliche Funktionen realisiert (8 Pkt)", "Auto / Manuell: unser Zusatz-Feature")
chain = [("Handy", "MQTT-App"), ("Node-RED", "validiert"), ("set_control.py", "schreibt Datei"), ("control_state.json", "mode + Werte")]
cw2, g2 = 2.6, 0.47
for i, (t, sub) in enumerate(chain):
    x = 0.75 + i * (cw2 + g2)
    bp.diagram_box(s, x, y + 0.1, cw2, 0.95, t, [sub], P)
    if i < 3:
        bp.diagram_arrow_h(s, x + cw2 + 0.04, x + cw2 + g2 - 0.04, y + 0.57, COL["p"])
bp.diagram_arrow_v(s, 0.75 + 3 * (cw2 + g2) + cw2 / 2, y + 1.05, y + 1.55, COL["p"], label="alle 5 s gelesen", label_side="left")
bp.diagram_box(s, 6.5, y + 1.6, 6.08, 0.95, "app.py (Backend aus Challenge I)", ["prüft bei jedem Messzyklus den Modus"], P, filled=True)
card(s, 0.75, y + 2.85, 5.75, 1.75, "auto",
     "Regellogik (`rules.py`) entscheidet anhand der **Temperatur**, wie stark Lüfter und Ventil laufen", size=13)
card(s, 6.83, y + 2.85, 5.75, 1.75, "manual  (Override)",
     "Pi nimmt die **Werte vom Handy** und schickt sie per I2C an den Aktor-Arduino", size=13)
bp.add_text(s, 0.75, y + 4.7, 11.8, 0.35, "Die automatische Regelung bleibt erhalten – wir haben nur einen zweiten Weg ergänzt.",
            size=13, color=bp.MUTED, italic=True)
notes(s, "Damit haben wir eine Zusatzfunktion gebaut: einen Umschalter zwischen automatisch und manuell. Das "
         "Python-Skript schreibt den Befehl vom Handy in eine kleine Datei, control_state.json. Unser Backend aus "
         "Challenge I schaut bei jedem Messzyklus, also alle fünf Sekunden, in diese Datei. Steht dort ‚auto‘, "
         "rechnet wie bisher die Regellogik aus, wie stark Lüfter und Ventil laufen. Steht dort ‚manual‘, nimmt "
         "das Backend die Werte vom Handy und schickt sie per I2C an den Aktor-Arduino. So kann man im Notfall "
         "per Hand eingreifen, zum Beispiel wenn ein Sensor spinnt. Die automatische Regelung bleibt dabei "
         "erhalten, wir haben nur einen zweiten Weg ergänzt.")

# ---------- Stand ----------
s, y = blank_content("Challenge II", "Stand", "Getestet & Stand")
items = [
    ("✓", "Flow läuft auf dem Pi", "Richtung Pi → Handy mit echten Live-Daten geprüft"),
    ("✓", "16 / 16 Tests grün", "Backend-Logik (Kalibrierung, Regeln, Datenbank, Steuerzustand) mit pytest"),
    ("✓", "Simulierter I2C-Bus", "Testen ohne angeschlossene Hardware möglich"),
    ("…", "Handy → Lüfter auf echter Hardware", "[VOR DEM VORTRAG TESTEN – dann ✓ oder ehrlich offen lassen]"),
]
for i, (mark, t, d) in enumerate(items):
    yy = y + i * 1.12
    done = mark == "✓"
    bp.add_oval(s, 0.75, yy, 0.62, 0.62, fill=COL["p"] if done else COL["light"], line=None if done else COL["p"])
    bp.add_text(s, 0.75, yy + 0.1, 0.62, 0.42, mark, size=20, color=bp.WHITE if done else COL["p"], bold=True, align=PP_ALIGN.CENTER)
    bp.add_text(s, 1.6, yy + 0.0, 10.9, 0.35, t, size=17, color=bp.FG, bold=True)
    bp.add_text(s, 1.6, yy + 0.4, 10.9, 0.35, d, size=13.5, color=bp.MUTED)
notes(s, "Zum Stand: Der Flow läuft auf dem Pi, und die Richtung vom Pi zum Handy haben wir mit echten Messdaten "
         "geprüft. Die Backend-Logik ist mit 16 automatischen Tests abgedeckt, die alle grün sind. Dafür haben "
         "wir den I2C-Bus simuliert, damit wir auch ohne angeschlossene Hardware testen können. Felix zeigt "
         "jetzt, wie das Ganze auf dem Handy aussieht.\n\n[Letzten Punkt vorher mit Felix klären und auf der Folie anpassen.]")

assert N == bp.TOTAL, N
bp.prs.save("../anton-folien.pptx")
print("Saved", N, "slides -> ../anton-folien.pptx")
